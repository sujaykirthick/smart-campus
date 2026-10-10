import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # blank layout

    # Institutional Color Palette
    BRAND_RED = RGBColor(164, 22, 26)       # #A4161A Deep institutional red
    TEXT_DARK = RGBColor(23, 23, 23)        # #171717 Near-black body text
    TEXT_MUTED = RGBColor(100, 100, 100)    # #646464 Muted text
    BG_LIGHT = RGBColor(245, 245, 245)      # #F5F5F5 Grey container
    BORDER_GREY = RGBColor(215, 215, 215)   # #D7D7D7 Border line
    CARD_BG = RGBColor(250, 250, 250)       # #FAFAFA Light card surface
    WHITE = RGBColor(255, 255, 255)

    FONT_NAME = "Calibri"

    def add_header(slide, title_text, slide_num_str, category_text="AMRITA EVENTS • GENESIS'26"):
        # Category / Super-title
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(10.0), Inches(0.32))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = FONT_NAME
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = BRAND_RED

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.74), Inches(10.5), Inches(0.75))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = FONT_NAME
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = BRAND_RED

        # Slide Number
        num_box = slide.shapes.add_textbox(Inches(11.333), Inches(0.55), Inches(1.2), Inches(0.4))
        tf_num = num_box.text_frame
        tf_num.word_wrap = False
        tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.text = slide_num_str
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.font.name = FONT_NAME
        p_num.font.size = Pt(14)
        p_num.font.bold = True
        p_num.font.color.rgb = TEXT_MUTED

        # Thin header accent line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.58), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_GREY
        line.line.color.rgb = BORDER_GREY

    def add_screenshot_placeholder(slide, left, top, width, height, label):
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = BG_LIGHT
        box.line.color.rgb = BORDER_GREY
        box.line.width = Pt(1.5)
        
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = f"[Screenshot: {label}]"
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_NAME
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED
        return box

    def set_speaker_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    
    # Left vertical accent bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(0.12), Inches(4.7))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRAND_RED
    bar.line.fill.background()

    # Title & Subtitle box
    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.4), Inches(11.3), Inches(3.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0
    
    p_tag = tf1.paragraphs[0]
    p_tag.text = "GENESIS'26 REVIEW"
    p_tag.font.name = FONT_NAME
    p_tag.font.size = Pt(14)
    p_tag.font.bold = True
    p_tag.font.color.rgb = BRAND_RED
    p_tag.space_after = Pt(14)

    p_main = tf1.add_paragraph()
    p_main.text = "Amrita Events: One Verified Platform for Campus Events, Halls and Attendance"
    p_main.font.name = FONT_NAME
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = BRAND_RED
    p_main.space_after = Pt(18)

    p_sub = tf1.add_paragraph()
    p_sub.text = "A unified institutional portal replacing fragmented announcements, double-booked auditoriums, and paper-based on-duty attendance."
    p_sub.font.name = FONT_NAME
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = TEXT_DARK

    # Metadata card at bottom
    meta_box = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(4.9), Inches(11.3), Inches(1.2))
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = CARD_BG
    meta_box.line.color.rgb = BORDER_GREY
    meta_box.line.width = Pt(1)
    
    tf_meta = meta_box.text_frame
    tf_meta.word_wrap = True
    tf_meta.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_meta.margin_left = Inches(0.3)
    
    p_m1 = tf_meta.paragraphs[0]
    p_m1.text = "Team: [Team Name / Number]"
    p_m1.font.name = FONT_NAME
    p_m1.font.size = Pt(18)
    p_m1.font.bold = True
    p_m1.font.color.rgb = BRAND_RED
    
    p_m2 = tf_meta.add_paragraph()
    p_m2.text = "Members: [Member 1, Member 2, Member 3, Member 4]  •  Amrita Vishwa Vidyapeetham"
    p_m2.font.name = FONT_NAME
    p_m2.font.size = Pt(16)
    p_m2.font.color.rgb = TEXT_MUTED

    set_speaker_notes(s1, "Introduction to Amrita Events (Genesis), an integrated institutional platform for event approvals, clash-free hall bookings, and verified digital attendance.")

    # =========================================================================
    # SLIDE 2: The Problem
    # =========================================================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_header(s2, "The Problem: Fragmented Campus Operations", "02 / 10")

    content_box2 = s2.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.733), Inches(5.0))
    tf2 = content_box2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0

    bullets_s2 = [
        "Fragmented Channels: Events are announced across disparate WhatsApp groups and Instagram stories, causing students to miss key deadlines.",
        "Venue Double-Booking: Multiple clubs frequently reserve the same hall or auditorium due to lack of a real-time availability calendar.",
        "Unverified OD Process: On-Duty (OD) letters are chased via manual paper signatures and unverified spreadsheets, leading to delays and proxy attendance.",
        "Redundant Coordination: Organizers repeatedly recreate proposal forms, volunteer rosters, and check-in sheets for every single event.",
        "[Survey result: X% of N students missed a major campus event due to fragmented communication]"
    ]

    for i, b_text in enumerate(bullets_s2):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = "•  " + b_text
        p.font.name = FONT_NAME
        p.font.size = Pt(21)
        p.font.color.rgb = BRAND_RED if i == 4 else TEXT_DARK
        if i == 4:
            p.font.bold = True
        p.space_after = Pt(18)

    set_speaker_notes(s2, "Campus event workflows currently suffer from fragmented communication, double-booked halls, and manual, unverifiable OD attendance chasing.")

    # =========================================================================
    # SLIDE 3: Our Solution (One Flow)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_header(s3, "Our Solution: One Unified Event Lifecycle", "03 / 10")

    steps = [
        ("1. Propose", "Club submits event\n& requests hall"),
        ("2. Approve", "Student Welfare\nreviews & approves"),
        ("3. Register", "Students register\n(waitlist if full)"),
        ("4. Check-In", "Door QR scan via\nauthorized phones"),
        ("5. Export OD", "Verified attendance\nexported to CSV"),
        ("6. Insights", "Gemini AI synthesizes\nreviews & metrics")
    ]

    step_w = Inches(1.75)
    step_h = Inches(3.2)
    step_y = Inches(2.1)
    gap = Inches(0.24)
    start_x = Inches(0.8)

    for i, (stitle, sdesc) in enumerate(steps):
        x = start_x + i * (step_w + gap)
        
        # Step card
        card = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, step_y, step_w, step_h)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BRAND_RED
        card.line.width = Pt(1.5)
        
        # Header banner inside card
        banner = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, step_y, step_w, Inches(0.65))
        banner.fill.solid()
        banner.fill.fore_color.rgb = BRAND_RED
        banner.line.fill.background()
        tf_b = banner.text_frame
        tf_b.word_wrap = True
        tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_b = tf_b.paragraphs[0]
        p_b.text = stitle
        p_b.alignment = PP_ALIGN.CENTER
        p_b.font.name = FONT_NAME
        p_b.font.size = Pt(17)
        p_b.font.bold = True
        p_b.font.color.rgb = WHITE

        # Body text inside card
        tb = s3.shapes.add_textbox(x + Inches(0.1), step_y + Inches(0.8), step_w - Inches(0.2), step_h - Inches(0.9))
        tf_desc = tb.text_frame
        tf_desc.word_wrap = True
        tf_desc.margin_left = tf_desc.margin_top = tf_desc.margin_right = tf_desc.margin_bottom = 0
        p_d = tf_desc.paragraphs[0]
        p_d.text = sdesc
        p_d.alignment = PP_ALIGN.CENTER
        p_d.font.name = FONT_NAME
        p_d.font.size = Pt(17)
        p_d.font.color.rgb = TEXT_DARK

        # Arrow between steps (except last)
        if i < len(steps) - 1:
            arrow_x = x + step_w + Inches(0.04)
            arrow_y = step_y + Inches(1.4)
            arrow = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arrow_x, arrow_y, Inches(0.16), Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = BRAND_RED
            arrow.line.fill.background()

    # Bottom summary box
    sum_box = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.7), Inches(11.733), Inches(0.95))
    sum_box.fill.solid()
    sum_box.fill.fore_color.rgb = CARD_BG
    sum_box.line.color.rgb = BORDER_GREY
    sum_box.line.width = Pt(1)
    tf_sum = sum_box.text_frame
    tf_sum.word_wrap = True
    tf_sum.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_sum.margin_left = Inches(0.3)
    p_sum = tf_sum.paragraphs[0]
    p_sum.text = "A single unbroken lifecycle: Zero paper forms, zero double-booked venues, and 100% verified on-duty records."
    p_sum.font.name = FONT_NAME
    p_sum.font.size = Pt(19)
    p_sum.font.bold = True
    p_sum.font.color.rgb = BRAND_RED

    set_speaker_notes(s3, "A unified, single-flow lifecycle from club proposal and clash-free approval to live QR check-in, verified OD export, and AI post-mortems.")

    # =========================================================================
    # SLIDE 4: Key Features
    # =========================================================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_header(s4, "Key Platform Features", "04 / 10")

    content_box4 = s4.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.733), Inches(5.0))
    tf4 = content_box4.text_frame
    tf4.word_wrap = True
    tf4.margin_left = tf4.margin_top = tf4.margin_right = tf4.margin_bottom = 0

    features = [
        "Public Clash-Free Hall Calendar: Real-time auditorium availability preventing scheduling conflicts before submission.",
        "Proposal & Approval Workflow: Departmental review with mandatory re-approval triggers whenever date, time, or venue changes.",
        "Fest Controls & Registration Engine: Dynamic fest toggles by Student Welfare with automated waitlist promotion on cancellations.",
        "Volunteer Management & QR Scanner: Organizer-approved volunteer scanning rosters with time-window-enforced camera check-ins.",
        "Verified OD Export & AI Insights: One-click CSV export of audited check-in logs paired with Gemini-generated event analytics."
    ]

    for i, f_text in enumerate(features):
        p = tf4.paragraphs[0] if i == 0 else tf4.add_paragraph()
        p.text = "•  " + f_text
        p.font.name = FONT_NAME
        p.font.size = Pt(21)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(18)

    set_speaker_notes(s4, "Core feature set covering hall clash prevention, fest controls, auto-waitlist promotion, verified digital passes, and feedback analytics.")

    # =========================================================================
    # SLIDE 5: Demo Walkthrough
    # =========================================================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_header(s5, "System Walkthrough & Interface", "05 / 10")

    cols = [
        ("Step 1: Club Proposes", "Club Submits Event & Hall Request", "Event Proposal & Clash Check"),
        ("Step 2: Welfare Approves", "Student Welfare Approves Venue", "Welfare Approval Queue"),
        ("Step 3: Student Registers", "Student Registers & Receives QR Pass", "Digital QR Pass & Waitlist"),
        ("Step 4: Scanner & OD", "Organizer Scans & Exports Verified OD", "Live Scanner & OD CSV Export")
    ]

    col_w = Inches(2.75)
    col_gap = Inches(0.24)
    start_x5 = Inches(0.8)

    for i, (step_tag, step_title, placeholder_lbl) in enumerate(cols):
        cx = start_x5 + i * (col_w + col_gap)
        
        # Step header box
        h_box = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, Inches(1.85), col_w, Inches(0.75))
        h_box.fill.solid()
        h_box.fill.fore_color.rgb = CARD_BG
        h_box.line.color.rgb = BRAND_RED
        h_box.line.width = Pt(1)
        tf_h = h_box.text_frame
        tf_h.word_wrap = True
        tf_h.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_ht = tf_h.paragraphs[0]
        p_ht.text = step_tag
        p_ht.alignment = PP_ALIGN.CENTER
        p_ht.font.name = FONT_NAME
        p_ht.font.size = Pt(14)
        p_ht.font.bold = True
        p_ht.font.color.rgb = BRAND_RED
        
        p_hd = tf_h.add_paragraph()
        p_hd.text = step_title
        p_hd.alignment = PP_ALIGN.CENTER
        p_hd.font.name = FONT_NAME
        p_hd.font.size = Pt(11)
        p_hd.font.color.rgb = TEXT_DARK

        # Screenshot placeholder
        add_screenshot_placeholder(s5, cx, Inches(2.7), col_w, Inches(4.2), placeholder_lbl)

    set_speaker_notes(s5, "A walkthrough of the 4 key stages: proposal submission, admin approval, student QR pass issuance, and door scanning with instant OD generation.")

    # =========================================================================
    # SLIDE 6: Tech Stack
    # =========================================================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_header(s6, "Production Architecture & Tech Stack", "06 / 10")

    tech_items = [
        ("Frontend Application", "Next.js 16 (App Router, Turbopack, TypeScript), Tailwind CSS, Lucide Icons"),
        ("Backend & Database", "Supabase (PostgreSQL 15, Auth, Storage, Serverless Functions)"),
        ("Database Security", "Row Level Security (RLS) policies and PostgreSQL exclusion triggers"),
        ("QR Pass & Scanner", "qrcode.react (SVG generation) and html5-qrcode (live camera feed scanner)"),
        ("AI Intelligence", "Google Gemini API (gemini-2.5-flash) with server-side caching & retry backoff"),
        ("Cloud Infrastructure", "Vercel Edge Platform and Supabase Managed Cloud Database")
    ]

    # 2 columns x 3 rows grid
    grid_w = Inches(5.7)
    grid_h = Inches(1.4)
    gx_start = Inches(0.8)
    gy_start = Inches(1.9)
    g_gap_x = Inches(0.33)
    g_gap_y = Inches(0.25)

    for i, (layer_title, layer_desc) in enumerate(tech_items):
        col_idx = i % 2
        row_idx = i // 2
        
        box_x = gx_start + col_idx * (grid_w + g_gap_x)
        box_y = gy_start + row_idx * (grid_h + g_gap_y)
        
        tbox = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, box_x, box_y, grid_w, grid_h)
        tbox.fill.solid()
        tbox.fill.fore_color.rgb = CARD_BG
        tbox.line.color.rgb = BORDER_GREY
        tbox.line.width = Pt(1)
        
        tf_t = tbox.text_frame
        tf_t.word_wrap = True
        tf_t.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_t.margin_left = Inches(0.25)
        tf_t.margin_right = Inches(0.25)
        
        p_t1 = tf_t.paragraphs[0]
        p_t1.text = layer_title
        p_t1.font.name = FONT_NAME
        p_t1.font.size = Pt(20)
        p_t1.font.bold = True
        p_t1.font.color.rgb = BRAND_RED
        p_t1.space_after = Pt(4)
        
        p_t2 = tf_t.add_paragraph()
        p_t2.text = layer_desc
        p_t2.font.name = FONT_NAME
        p_t2.font.size = Pt(17)
        p_t2.font.color.rgb = TEXT_DARK

    set_speaker_notes(s6, "Built on a full-stack Next.js and Supabase architecture, leveraging PostgreSQL exclusion constraints and server-side Gemini 2.5 Flash.")

    # =========================================================================
    # SLIDE 7: Engineering Decisions
    # =========================================================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    add_header(s7, "Engineering Decisions: Integrity & Enforcement", "07 / 10")

    content_box7 = s7.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.733), Inches(5.0))
    tf7 = content_box7.text_frame
    tf7.word_wrap = True
    tf7.margin_left = tf7.margin_top = tf7.margin_right = tf7.margin_bottom = 0

    decisions = [
        "Postgres Exclusion Constraints: Venue collisions are physically prevented at the database engine level, making double-approvals impossible.",
        "Trigger-Enforced State Machine: Approval status transitions and re-approval resets on schedule edits are strictly enforced via DB triggers.",
        "Cryptographic QR Security: Random 24-byte tokens validated server-side, restricted to authorized coordinators within an active time window.",
        "Atomic Waitlist Promotion: Database-level transactions ensure immediate, race-condition-free promotion to confirmed status upon cancellation.",
        "Deterministic Roll Parsing: Student department, batch, and section are derived directly from validated university roll number patterns."
    ]

    for i, d_text in enumerate(decisions):
        p = tf7.paragraphs[0] if i == 0 else tf7.add_paragraph()
        p.text = "•  " + d_text
        p.font.name = FONT_NAME
        p.font.size = Pt(21)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(18)

    set_speaker_notes(s7, "Engineered with database-level exclusion constraints for zero hall clashes, atomic waitlist promotion, and cryptographically verified check-ins.")

    # =========================================================================
    # SLIDE 8: AI Usage
    # =========================================================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    add_header(s8, "AI Integration: Grounded Feedback Synthesis", "08 / 10")

    # Left column: Bullets (Width: 6.8 inches)
    content_box8 = s8.shapes.add_textbox(Inches(0.8), Inches(1.85), Inches(6.8), Inches(5.1))
    tf8 = content_box8.text_frame
    tf8.word_wrap = True
    tf8.margin_left = tf8.margin_top = tf8.margin_right = tf8.margin_bottom = 0

    ai_bullets = [
        "Qualitative Synthesis: Gemini extracts recurring themes, operational strengths, and bottlenecks from student reviews.",
        "SQL-Grounded Advice: Advice derives directly from computed database metrics (turnout, capacity, no-show rate), not guesses.",
        "Verified Reviews Only: Only attendees marked checked-in by door scanners can submit reviews, eliminating fake feedback.",
        "Strict Prompt Isolation: Review comments are treated strictly as untrusted evaluation data, never as executable instructions.",
        "Auditable Verification: AI summaries are cached and displayed right next to raw attendance numbers for auditability."
    ]

    for i, b_text in enumerate(ai_bullets):
        p = tf8.paragraphs[0] if i == 0 else tf8.add_paragraph()
        p.text = "•  " + b_text
        p.font.name = FONT_NAME
        p.font.size = Pt(20)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(12)

    # Right column: Screenshot placeholder (Width: 4.633 inches)
    add_screenshot_placeholder(s8, Inches(7.9), Inches(1.85), Inches(4.633), Inches(5.0), "Event Insights & AI Synthesis")

    set_speaker_notes(s8, "Gemini AI analyzes verified attendee reviews and historical attendance rates to provide grounded post-mortems and overbooking recommendations.")

    # =========================================================================
    # SLIDE 9: Campus Digital Twin
    # =========================================================================
    s9 = prs.slides.add_slide(blank_slide_layout)
    add_header(s9, "Campus Digital Twin: Spatial Integration", "09 / 10")

    # Status Pill / Banner
    status_pill = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(6.8), Inches(0.48))
    status_pill.fill.solid()
    status_pill.fill.fore_color.rgb = CARD_BG
    status_pill.line.color.rgb = BRAND_RED
    status_pill.line.width = Pt(1)
    tf_sp = status_pill.text_frame
    tf_sp.word_wrap = True
    tf_sp.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_sp.margin_left = Inches(0.2)
    p_sp = tf_sp.paragraphs[0]
    p_sp.text = "Status: Working Integration  |  Planned: Indoor 3D Turn-by-Turn"
    p_sp.font.name = FONT_NAME
    p_sp.font.size = Pt(15)
    p_sp.font.bold = True
    p_sp.font.color.rgb = BRAND_RED

    # Left column: Bullets
    content_box9 = s9.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(6.8), Inches(4.4))
    tf9 = content_box9.text_frame
    tf9.word_wrap = True
    tf9.margin_left = tf9.margin_top = tf9.margin_right = tf9.margin_bottom = 0

    twin_bullets = [
        "Spatial Campus Mapping: Digital model representing campus blocks, auditoriums, and seminar halls.",
        "One-Tap Venue Pinning: Registered students tap 'Locate Venue' directly from their digital pass to view hall coordinates.",
        "Frictionless Wayfinding: Helps first-year students and external fest attendees locate venues without confusion.",
        "[Working: Live venue coordinate pin  •  Planned: Multi-floor indoor path routing]"
    ]

    for i, b_text in enumerate(twin_bullets):
        p = tf9.paragraphs[0] if i == 0 else tf9.add_paragraph()
        p.text = "•  " + b_text
        p.font.name = FONT_NAME
        p.font.size = Pt(21)
        p.font.color.rgb = BRAND_RED if i == 3 else TEXT_DARK
        if i == 3:
            p.font.bold = True
        p.space_after = Pt(18)

    # Right column: Screenshot placeholder
    add_screenshot_placeholder(s9, Inches(7.9), Inches(1.85), Inches(4.633), Inches(5.0), "Campus Digital Twin & Venue Locator")

    set_speaker_notes(s9, "Integration with the campus digital twin allows students to locate and navigate to booked auditoriums directly from their digital event pass.")

    # =========================================================================
    # SLIDE 10: Current Status & Next Steps
    # =========================================================================
    s10 = prs.slides.add_slide(blank_slide_layout)
    add_header(s10, "Current Status & Future Roadmap", "10 / 10")

    # Left Column: Built & Running Today
    left_card = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), Inches(5.7), Inches(4.0))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = WHITE
    left_card.line.color.rgb = BORDER_GREY
    left_card.line.width = Pt(1)
    
    # Header for left card
    l_hdr = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), Inches(5.7), Inches(0.55))
    l_hdr.fill.solid()
    l_hdr.fill.fore_color.rgb = CARD_BG
    l_hdr.line.color.rgb = BORDER_GREY
    tf_lh = l_hdr.text_frame
    tf_lh.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_lh.margin_left = Inches(0.2)
    p_lh = tf_lh.paragraphs[0]
    p_lh.text = "BUILT & RUNNING LIVE"
    p_lh.font.name = FONT_NAME
    p_lh.font.size = Pt(16)
    p_lh.font.bold = True
    p_lh.font.color.rgb = BRAND_RED

    tb_l = s10.shapes.add_textbox(Inches(0.95), Inches(2.5), Inches(5.4), Inches(3.2))
    tf_lb = tb_l.text_frame
    tf_lb.word_wrap = True
    tf_lb.margin_left = tf_lb.margin_top = tf_lb.margin_right = tf_lb.margin_bottom = 0
    
    built_items = [
        "End-to-end event workflow running in production",
        "Clash-free hall booking & Student Welfare approval",
        "QR pass check-in & instant OD CSV export",
        "Gemini AI post-event qualitative synthesis",
        "[Honest limit: Seeded demo data & demo accounts]"
    ]
    for i, item in enumerate(built_items):
        p = tf_lb.paragraphs[0] if i == 0 else tf_lb.add_paragraph()
        p.text = "•  " + item
        p.font.name = FONT_NAME
        p.font.size = Pt(18)
        p.font.color.rgb = BRAND_RED if i == 4 else TEXT_DARK
        p.space_after = Pt(10)

    # Right Column: Next Steps
    right_card = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.833), Inches(1.85), Inches(5.7), Inches(4.0))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = WHITE
    right_card.line.color.rgb = BORDER_GREY
    right_card.line.width = Pt(1)

    r_hdr = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.833), Inches(1.85), Inches(5.7), Inches(0.55))
    r_hdr.fill.solid()
    r_hdr.fill.fore_color.rgb = CARD_BG
    r_hdr.line.color.rgb = BORDER_GREY
    tf_rh = r_hdr.text_frame
    tf_rh.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_rh.margin_left = Inches(0.2)
    p_rh = tf_rh.paragraphs[0]
    p_rh.text = "ROADMAP & NEXT STEPS"
    p_rh.font.name = FONT_NAME
    p_rh.font.size = Pt(16)
    p_rh.font.bold = True
    p_rh.font.color.rgb = BRAND_RED

    tb_r = s10.shapes.add_textbox(Inches(7.0), Inches(2.5), Inches(5.4), Inches(3.2))
    tf_rb = tb_r.text_frame
    tf_rb.word_wrap = True
    tf_rb.margin_left = tf_rb.margin_top = tf_rb.margin_right = tf_rb.margin_bottom = 0

    roadmap_items = [
        "University Microsoft Outlook SSO integration",
        "Direct integration with central student records DB",
        "Automated push & SMS waitlist notifications",
        "Automated department-wise OD reporting to faculty",
        "Indoor multi-level digital twin navigation"
    ]
    for i, item in enumerate(roadmap_items):
        p = tf_rb.paragraphs[0] if i == 0 else tf_rb.add_paragraph()
        p.text = "•  " + item
        p.font.name = FONT_NAME
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(10)

    # Bottom Contact & Demo Bar
    footer_bar = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.8))
    footer_bar.fill.solid()
    footer_bar.fill.fore_color.rgb = CARD_BG
    footer_bar.line.color.rgb = BRAND_RED
    footer_bar.line.width = Pt(1.5)
    tf_foot = footer_bar.text_frame
    tf_foot.word_wrap = True
    tf_foot.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_foot.margin_left = Inches(0.3)
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "Thank You  •  Questions?  •  Live Demo: [https://genesis-amrita.vercel.app]"
    p_foot.font.name = FONT_NAME
    p_foot.font.size = Pt(20)
    p_foot.font.bold = True
    p_foot.font.color.rgb = BRAND_RED

    set_speaker_notes(s10, "The full core system is live and verified; our next phase focuses on institutional SSO integration, push notifications, and automated faculty OD dispatch.")

    # Save presentation
    os.makedirs("presentation", exist_ok=True)
    out_path = os.path.join("presentation", "amrita-events.pptx")
    prs.save(out_path)
    print(f"Presentation updated successfully at: {os.path.abspath(out_path)}")

if __name__ == "__main__":
    create_presentation()
